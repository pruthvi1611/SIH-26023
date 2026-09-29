"""
Multi-Format Export Engine for Phase 5: Automated Geological Reporting
SIH 2026 Problem Statement 26023 (CMPDI / Coal India Limited)

Provides professional export handlers for:
1. PDF (ReportLab - high-density professional layout with NumberedCanvas & CMPDI branding)
2. Excel (openpyxl - multi-sheet styled workbook with metadata & provenance)
3. Markdown (GitHub-flavored formatted text)
"""

import io
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional

# ReportLab imports for PDF generation
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
    HRFlowable,
)
from reportlab.pdfgen import canvas

# openpyxl imports for Excel generation
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.schemas.report import GeneratedReport, ReportSection, ReportProvenanceItem


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to dynamically compute and stamp 'Page X of Y' footers."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Running Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(40, 810, "Central Mine Planning & Design Institute Limited (CMPDI) — GeoMine Intelligence")
            self.drawRightString(555, 810, "CONFIDENTIAL / STATUTORY")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(40, 804, 555, 804)

        # Running Footer (all pages)
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(40, 38, 555, 38)

        footer_text = "Prepared by GeoMine Intelligence Platform | CIL / CMPDI Exploration Standard"
        self.drawString(40, 26, footer_text)
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(555, 26, page_str)

        self.restoreState()


class ReportExporter:
    """Handles deterministic file exports into PDF, Excel, and Markdown."""

    @staticmethod
    def export_pdf(report: GeneratedReport) -> bytes:
        """Generates a professional CMPDI geological report PDF."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=40,
            rightMargin=40,
            topMargin=45,
            bottomMargin=50,
        )

        styles = getSampleStyleSheet()

        # Custom Palette
        c_primary = colors.HexColor("#0B192C")      # Deep navy
        c_secondary = colors.HexColor("#1E293B")    # Slate dark
        c_accent = colors.HexColor("#D97706")       # Amber
        c_text = colors.HexColor("#0F172A")         # Dark charcoal
        c_muted = colors.HexColor("#64748B")        # Slate light
        c_border = colors.HexColor("#CBD5E1")       # Border grey
        c_header_bg = colors.HexColor("#0F172A")    # Table header dark
        c_row_alt = colors.HexColor("#F8FAFC")      # Alternating row

        # Custom Typography
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=c_primary,
            spaceAfter=4,
        )

        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=13,
            textColor=c_accent,
            spaceAfter=8,
        )

        meta_label_style = ParagraphStyle(
            "MetaLabel",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=c_secondary,
        )

        meta_val_style = ParagraphStyle(
            "MetaVal",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=c_text,
        )

        h2_style = ParagraphStyle(
            "SectionH2",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=c_secondary,
            spaceBefore=14,
            spaceAfter=6,
            keepWithNext=True,
        )

        summary_box_style = ParagraphStyle(
            "SummaryBoxText",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=8.5,
            leading=11,
            textColor=c_text,
        )

        th_style = ParagraphStyle(
            "TableHead",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7.5,
            leading=9,
            textColor=colors.white,
            alignment=1,  # Center
        )

        td_style = ParagraphStyle(
            "TableCell",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.5,
            textColor=c_text,
        )

        td_num_style = ParagraphStyle(
            "TableCellNum",
            parent=styles["Normal"],
            fontName="Courier",
            fontSize=7.5,
            leading=9.5,
            textColor=c_text,
            alignment=2,  # Right
        )

        provenance_style = ParagraphStyle(
            "ProvEvidence",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=7,
            leading=8.5,
            textColor=c_muted,
        )

        elements = []

        # 1. Header Banner & Title
        elements.append(Paragraph(report.metadata.title, title_style))
        elements.append(Paragraph(f"{report.metadata.organization.upper()} — EXPLORATION DIVISION", subtitle_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=c_accent, spaceBefore=2, spaceAfter=8))

        # 2. Metadata Grid
        meta_table_data = [
            [
                Paragraph("<b>Operating Subsidiary:</b>", meta_label_style),
                Paragraph(report.metadata.subsidiary or "Coal India Limited", meta_val_style),
                Paragraph("<b>Date Generated:</b>", meta_label_style),
                Paragraph(report.metadata.generated_at.strftime("%Y-%m-%d %H:%M UTC"), meta_val_style),
            ],
            [
                Paragraph("<b>Data Scope:</b>", meta_label_style),
                Paragraph(f"{report.metadata.scope.value.upper()} ({report.metadata.scope_target})", meta_val_style),
                Paragraph("<b>Report Identifier:</b>", meta_label_style),
                Paragraph(report.metadata.report_id[:18] + "...", meta_val_style),
            ],
            [
                Paragraph("<b>Records Consumed:</b>", meta_label_style),
                Paragraph(f"{report.metadata.total_records_used} Entities", meta_val_style),
                Paragraph("<b>Issuing Authority:</b>", meta_label_style),
                Paragraph("CMPDI / GeoMine Intelligence", meta_val_style),
            ],
        ]
        meta_table = Table(meta_table_data, colWidths=[110, 150, 100, 155])
        meta_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), c_row_alt),
            ("BOX", (0, 0), (-1, -1), 0.5, c_border),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(meta_table)
        elements.append(Spacer(1, 8))

        # 3. Executive Highlights Strip
        stats = report.summary_statistics
        stat_boxes = []
        if stats.get("total_gross_reserves_mt") is not None:
            stat_boxes.append(f"Gross Reserves: {stats['total_gross_reserves_mt']:,} MT")
        if stats.get("total_extractable_reserves_mt") is not None:
            stat_boxes.append(f"Extractable: {stats['total_extractable_reserves_mt']:,} MT")
        if stats.get("overall_recovery_percentage") is not None:
            stat_boxes.append(f"Recovery: {stats['overall_recovery_percentage']}%")
        if stats.get("total_boreholes_logged"):
            stat_boxes.append(f"Boreholes: {stats['total_boreholes_logged']}")
        if stats.get("total_seams_analyzed"):
            stat_boxes.append(f"Seams: {stats['total_seams_analyzed']}")
        if stats.get("average_ash_percentage") is not None:
            stat_boxes.append(f"Avg Ash: {stats['average_ash_percentage']}%")
        if stats.get("average_gcv_kcal_kg") is not None:
            stat_boxes.append(f"Avg GCV: {stats['average_gcv_kcal_kg']:,} kcal/kg")

        if stat_boxes:
            highlight_text = " | ".join(stat_boxes)
            p_high = Paragraph(f"<b>Key Highlights:</b> {highlight_text}", summary_box_style)
            high_table = Table([[p_high]], colWidths=[515])
            high_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FEF3C7")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#F59E0B")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]))
            elements.append(high_table)
            elements.append(Spacer(1, 6))

        # 4. Render Structured Report Sections
        usable_width = 515
        for sec in report.sections:
            sec_elements = []
            sec_elements.append(Paragraph(sec.title, h2_style))

            if sec.summary_text:
                sec_elements.append(Paragraph(f"<i>{sec.summary_text}</i>", summary_box_style))
                sec_elements.append(Spacer(1, 4))

            if sec.headers and sec.rows:
                num_cols = len(sec.headers)
                col_w = usable_width / num_cols
                # Specific customized column widths if needed
                if num_cols == 7:
                    col_widths = [80, 80, 70, 65, 75, 65, 80]
                elif num_cols == 6:
                    col_widths = [90, 85, 75, 75, 100, 90]
                elif num_cols == 10:
                    col_widths = [60, 50, 45, 45, 45, 45, 55, 55, 55, 60]
                elif num_cols == 12:
                    col_widths = [50, 40, 35, 35, 35, 35, 45, 50, 50, 45, 45, 50]
                else:
                    col_widths = [col_w] * num_cols

                t_data = []
                # Header row
                header_row = [Paragraph(f"<b>{h}</b>", th_style) for h in sec.headers]
                t_data.append(header_row)

                # Data rows
                for r_idx, row in enumerate(sec.rows):
                    row_cells = []
                    for c_idx, cell in enumerate(row):
                        cell_str = str(cell) if cell is not None else "N/A"
                        # Right-align numeric columns
                        if any(c in cell_str for c in ["MT", "%", "m", "Yrs"]) or cell_str.replace(".", "", 1).isdigit():
                            row_cells.append(Paragraph(cell_str, td_num_style))
                        else:
                            row_cells.append(Paragraph(cell_str, td_style))
                    t_data.append(row_cells)

                t = Table(t_data, colWidths=col_widths, repeatRows=1)
                t_style = [
                    ("BACKGROUND", (0, 0), (-1, 0), c_header_bg),
                    ("BOX", (0, 0), (-1, -1), 0.5, c_border),
                    ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ]
                # Alternating row colors
                for r_idx in range(1, len(t_data)):
                    if r_idx % 2 == 0:
                        t_style.append(("BACKGROUND", (0, r_idx), (-1, r_idx), c_row_alt))
                t.setStyle(TableStyle(t_style))
                sec_elements.append(t)
            else:
                sec_elements.append(Paragraph("<i>No data records available for this section in the current scope.</i>", summary_box_style))

            sec_elements.append(Spacer(1, 8))
            if len(sec.rows or []) <= 6:
                elements.append(KeepTogether(sec_elements))
            else:
                elements.extend(sec_elements)

        # 5. Provenance & Statutory Audit Trail
        prov_elements = []
        prov_elements.append(Paragraph("Statutory Provenance & Audit Trail", h2_style))
        prov_elements.append(
            Paragraph("Every factual figure presented in this report is 100% grounded in verified source documents:", summary_box_style)
        )
        prov_elements.append(Spacer(1, 4))

        if report.provenance_sources:
            prov_headers = ["Source Document", "Page", "Category", "Verified Verbatim Evidence"]
            prov_data = [[Paragraph(f"<b>{h}</b>", th_style) for h in prov_headers]]
            for p in report.provenance_sources[:40]:  # Cap to top 40 citations in print
                clean_ev = p.evidence_text.replace("\n", " ").strip()
                if len(clean_ev) > 120:
                    clean_ev = clean_ev[:117] + "..."
                prov_data.append([
                    Paragraph(p.source_document, td_style),
                    Paragraph(str(p.source_page), td_num_style),
                    Paragraph(p.entity_type, td_style),
                    Paragraph(f'"{clean_ev}"', provenance_style),
                ])

            pt = Table(prov_data, colWidths=[120, 35, 80, 280], repeatRows=1)
            pt_style = [
                ("BACKGROUND", (0, 0), (-1, 0), c_secondary),
                ("BOX", (0, 0), (-1, -1), 0.5, c_border),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
            for r_idx in range(1, len(prov_data)):
                if r_idx % 2 == 0:
                    pt_style.append(("BACKGROUND", (0, r_idx), (-1, r_idx), c_row_alt))
            pt.setStyle(TableStyle(pt_style))
            prov_elements.append(pt)
        else:
            prov_elements.append(Paragraph("<i>No provenance citations registered.</i>", summary_box_style))

        elements.extend(prov_elements)

        # Build PDF with two-pass canvas for total page count
        doc.build(elements, canvasmaker=NumberedCanvas)
        return buffer.getvalue()

    @staticmethod
    def export_excel(report: GeneratedReport) -> bytes:
        """Generates a structured, multi-worksheet Excel workbook with professional styling."""
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Common styles
        header_fill = PatternFill(start_color="0B192C", end_color="0B192C", fill_type="solid")
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="0B192C")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        sheet_name_map = {
            "mine_planning_overview": "Mines & Projects",
            "stratigraphic_borehole_data": "Boreholes",
            "coal_seam_stratigraphy": "Coal Seams",
            "proximate_analysis_summary": "Proximate Analysis",
            "geological_metrics_summary": "Geological Metrics",
            "seam_reserve_reconciliation_table": "Seam Reconciliation",
            "correlated_stratigraphy": "Stratigraphy",
            "reconciliation_quality_matrix": "Quality Matrix",
        }

        # 1. Summary Sheet
        ws_sum = wb.create_sheet(title="Summary")
        ws_sum.views.sheetView[0].showGridLines = True

        ws_sum["A1"] = report.metadata.title
        ws_sum["A1"].font = title_font
        ws_sum["A2"] = f"{report.metadata.organization} — {report.metadata.subsidiary or 'Coal India Limited'}"
        ws_sum["A2"].font = Font(name="Calibri", size=11, bold=True, color="D97706")

        meta_rows = [
            ("Report Identifier", report.metadata.report_id),
            ("Report Type", report.metadata.report_type.value.replace("_", " ").title()),
            ("Data Scope", f"{report.metadata.scope.value.upper()} ({report.metadata.scope_target})"),
            ("Generation Timestamp (UTC)", report.metadata.generated_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("Total Records Consumed", report.metadata.total_records_used),
        ]

        ws_sum["A4"] = "REPORT METADATA"
        ws_sum["A4"].font = bold_font
        for i, (k, v) in enumerate(meta_rows, start=5):
            ws_sum[f"A{i}"] = k
            ws_sum[f"A{i}"].font = bold_font
            ws_sum[f"B{i}"] = v
            ws_sum[f"B{i}"].font = regular_font

        ws_sum["A11"] = "EXECUTIVE SUMMARY STATISTICS"
        ws_sum["A11"].font = bold_font
        stat_rows = [
            ("Mines & Projects Evaluated", report.summary_statistics.get("total_mines_projects", 0)),
            ("Exploratory Boreholes Logged", report.summary_statistics.get("total_boreholes_logged", 0)),
            ("Coal Seams Analyzed", report.summary_statistics.get("total_seams_analyzed", 0)),
            ("Proximate Quality Analyses", report.summary_statistics.get("total_proximate_records", 0)),
            ("Geological Metrics Reported", report.summary_statistics.get("total_geological_metrics", 0)),
            ("Gross Reserves (MT)", report.summary_statistics.get("total_gross_reserves_mt") or "N/A"),
            ("Extractable Reserves (MT)", report.summary_statistics.get("total_extractable_reserves_mt") or "N/A"),
            ("Estimated Recovery Factor (%)", report.summary_statistics.get("overall_recovery_percentage") or "N/A"),
            ("Mean Seam True Thickness (m)", report.summary_statistics.get("mean_seam_thickness_m") or "N/A"),
            ("Average Ash Percentage (%)", report.summary_statistics.get("average_ash_percentage") or "N/A"),
            ("Average GCV (kcal/kg)", report.summary_statistics.get("average_gcv_kcal_kg") or "N/A"),
        ]
        for i, (k, v) in enumerate(stat_rows, start=12):
            ws_sum[f"A{i}"] = k
            ws_sum[f"A{i}"].font = bold_font
            ws_sum[f"B{i}"] = v
            ws_sum[f"B{i}"].font = regular_font

        ws_sum.column_dimensions["A"].width = 34
        ws_sum.column_dimensions["B"].width = 45

        # 2. Section Worksheets
        used_sheet_names = {"Summary"}
        for sec in report.sections:
            sheet_title = sheet_name_map.get(sec.section_id, sec.title.split(". ", 1)[-1])
            # Excel limits sheet names to 31 chars
            safe_title = sheet_title[:31].replace("/", "-").replace(":", "")
            if safe_title in used_sheet_names:
                suffix = 2
                while f"{safe_title[:28]}_{suffix}" in used_sheet_names:
                    suffix += 1
                safe_title = f"{safe_title[:28]}_{suffix}"
            used_sheet_names.add(safe_title)
            ws = wb.create_sheet(title=safe_title)
            ws.views.sheetView[0].showGridLines = True

            ws["A1"] = sec.title
            ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="0B192C")
            if sec.summary_text:
                ws["A2"] = sec.summary_text
                ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="475569")
                start_row = 4
            else:
                start_row = 3

            if sec.headers:
                for col_idx, header in enumerate(sec.headers, start=1):
                    cell = ws.cell(row=start_row, column=col_idx, value=header)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = thin_border
                ws.row_dimensions[start_row].height = 24

                if sec.rows:
                    for r_idx, row in enumerate(sec.rows, start=start_row + 1):
                        is_alt = (r_idx % 2 == 0)
                        for c_idx, val in enumerate(row, start=1):
                            cell = ws.cell(row=r_idx, column=c_idx, value=val)
                            cell.font = regular_font
                            cell.border = thin_border
                            if is_alt:
                                cell.fill = alt_fill
                            # Numeric alignment
                            if isinstance(val, (int, float)):
                                cell.alignment = Alignment(horizontal="right")
                            else:
                                cell.alignment = Alignment(horizontal="left")

                # Auto-adjust column width
                for col in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        if cell.row >= start_row and cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # 3. Provenance Audit Worksheet
        ws_prov = wb.create_sheet(title="Sources")
        ws_prov.views.sheetView[0].showGridLines = True
        ws_prov["A1"] = "Statutory Provenance & Audit Trail"
        ws_prov["A1"].font = Font(name="Calibri", size=12, bold=True, color="0B192C")
        ws_prov["A2"] = "Every extracted metric verified with source document citation and page number."
        ws_prov["A2"].font = Font(name="Calibri", size=9, italic=True, color="475569")

        prov_headers = ["Source Document", "Page", "Entity Category", "Entity Identifier", "Verbatim Evidence Excerpt"]
        for c_idx, h in enumerate(prov_headers, start=1):
            cell = ws_prov.cell(row=4, column=c_idx, value=h)
            cell.fill = sub_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for r_idx, p in enumerate(report.provenance_sources, start=5):
            is_alt = (r_idx % 2 == 0)
            row_vals = [p.source_document, p.source_page, p.entity_type, p.entity_id, p.evidence_text]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_prov.cell(row=r_idx, column=c_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                if is_alt:
                    cell.fill = alt_fill
                if c_idx == 2:
                    cell.alignment = Alignment(horizontal="right")

        ws_prov.column_dimensions["A"].width = 30
        ws_prov.column_dimensions["B"].width = 10
        ws_prov.column_dimensions["C"].width = 20
        ws_prov.column_dimensions["D"].width = 24
        ws_prov.column_dimensions["E"].width = 65

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def export_markdown(report: GeneratedReport) -> str:
        """Returns the pre-rendered markdown content."""
        return report.markdown_content


report_exporter = ReportExporter()
