"""
Document Ingestion & Multi-Modal Parser Service
Handles file validation, secure storage, and multi-format extraction:
- PDF: Native PyMuPDF (fitz) text, blocks, tables + OCR fallback for scanned pages
- DOCX: python-docx paragraphs and structured tables
- XLSX: openpyxl multi-sheet structured tables and cell grids
- Images: PIL + Tesseract OCR fallback
"""

import io
import re
import json
import uuid
import time
import hashlib
from pathlib import Path
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional, Tuple

import pymupdf
import docx
import openpyxl
from PIL import Image

from app.core.config import settings
from app.schemas.document import (
    DocumentMetadata,
    DocumentPage,
    TextBlock,
    ExtractedTable,
)
from app.services.ocr.service import OCRDocumentParserService
from app.services.documents.db import doc_db
from app.services.documents.table_normalization import (
    normalize_table_text,
    is_table_doubled,
)


class DocumentIngestionService:
    """Service for validating, storing, and parsing mining and geological documents."""

    ALLOWED_EXTENSIONS = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".xlsx": "xlsx",
        ".png": "png",
        ".jpg": "jpg",
        ".jpeg": "jpg",
        ".tiff": "tiff",
        ".tif": "tiff",
    }

    ALLOWED_MIME_TYPES = {
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "image/png",
        "image/jpeg",
        "image/tiff",
        "application/octet-stream",  # Permitted if extension is explicitly valid
    }

    MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50 MB

    def __init__(self, upload_dir: Optional[Path] = None, processed_dir: Optional[Path] = None):
        self.upload_dir = upload_dir or settings.UPLOAD_DIRECTORY
        self.processed_dir = processed_dir or settings.PROCESSED_DIRECTORY
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self.processed_dir.mkdir(parents=True, exist_ok=True)
        self.ocr_service = OCRDocumentParserService()

    def sanitize_filename(self, filename: str) -> str:
        """Removes path traversals and special characters, retaining extension."""
        clean = Path(filename).name
        clean = re.sub(r"[^a-zA-Z0-9_.-]", "_", clean)
        return clean or "document"

    def validate_file(self, filename: str, content_type: str, file_size: int) -> Tuple[bool, Optional[str], str]:
        """
        Validates file extension, MIME type, and size constraints.
        Returns (is_valid, error_message, file_type).
        """
        ext = Path(filename).suffix.lower()
        if ext not in self.ALLOWED_EXTENSIONS:
            allowed = ", ".join(self.ALLOWED_EXTENSIONS.keys())
            return False, f"Unsupported file extension '{ext}'. Allowed extensions: {allowed}", ""

        file_type = self.ALLOWED_EXTENSIONS[ext]

        if content_type and content_type not in self.ALLOWED_MIME_TYPES:
            # Allow fallback if extension is trusted
            if content_type not in ["application/octet-stream", "binary/octet-stream"]:
                return False, f"Unsupported MIME type '{content_type}'.", file_type

        if file_size > self.MAX_FILE_SIZE_BYTES:
            max_mb = self.MAX_FILE_SIZE_BYTES // (1024 * 1024)
            return False, f"File size exceeds maximum allowed limit of {max_mb}MB.", file_type

        if file_size <= 0:
            return False, "Uploaded file is empty (0 bytes).", file_type

        return True, None, file_type

    def _parse_pdf(self, file_path: Path) -> List[DocumentPage]:
        """
        Parses PDF using PyMuPDF (fitz) for native text, bounding blocks, and tables.
        Applies OCR fallback to image-heavy or scanned pages.
        """
        pages: List[DocumentPage] = []
        doc = pymupdf.open(file_path)

        try:
            for page_index in range(len(doc)):
                page = doc[page_index]
                page_num = page_index + 1
                rect = page.rect
                dim = {"width": round(rect.width, 2), "height": round(rect.height, 2)}

                # 1. Native text extraction
                native_text = page.get_text("text").strip()

                # 2. Extract bounding-box blocks
                blocks: List[TextBlock] = []
                raw_blocks = page.get_text("blocks")  # (x0, y0, x1, y1, text, block_no, block_type)
                for b in raw_blocks:
                    if len(b) >= 5 and b[4].strip():
                        blocks.append(
                            TextBlock(
                                block_index=b[5] if len(b) > 5 else len(blocks),
                                bbox=[round(coord, 2) for coord in b[:4]],
                                text=b[4].strip(),
                                block_type="text" if len(b) <= 6 or b[6] == 0 else "image",
                            )
                        )

                # 3. Extract tables using PyMuPDF table finder
                tables: List[ExtractedTable] = []
                try:
                    tab_finder = page.find_tables()
                    for t_idx, tab in enumerate(tab_finder.tables):
                        extracted_tab = tab.extract()
                        if extracted_tab and len(extracted_tab) > 0:
                            tab_doubled = is_table_doubled(extracted_tab)
                            headers = [
                                normalize_table_text(str(c or "").strip(), tab_doubled)
                                for c in extracted_tab[0]
                            ]
                            rows = [
                                [
                                    normalize_table_text(str(c or "").strip(), tab_doubled)
                                    for c in row
                                ]
                                for row in extracted_tab[1:]
                            ]
                            tables.append(
                                ExtractedTable(
                                    table_index=t_idx + 1,
                                    num_rows=len(extracted_tab),
                                    num_cols=len(headers),
                                    headers=headers,
                                    rows=rows,
                                    bbox=[round(c, 2) for c in tab.bbox] if hasattr(tab, "bbox") else None,
                                )
                            )
                except Exception:
                    pass  # Non-fatal if table detection encounters complex vector graphics

                # 4. Low-text / Scanned page detection & OCR fallback
                word_count = len(native_text.split())
                char_count = len(native_text)
                is_scanned_or_low_text = (word_count < 10 or char_count < 50) and len(page.get_images()) > 0

                ocr_applied = False
                ocr_conf = None
                ocr_error = None
                final_text = native_text

                if is_scanned_or_low_text:
                    try:
                        # Render page to PIL image
                        pix = page.get_pixmap(dpi=150)
                        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        ocr_text, ocr_conf, ocr_err, ocr_boxes = self.ocr_service.ocr_image(img)

                        if ocr_text:
                            ocr_applied = True
                            if final_text:
                                final_text = f"{final_text}\n\n[OCR Fallback Text]:\n{ocr_text}"
                            else:
                                final_text = ocr_text

                        if ocr_boxes and not blocks:
                            for b in ocr_boxes:
                                blocks.append(
                                    TextBlock(
                                        block_index=b["index"],
                                        bbox=b["bbox"],
                                        text=b["text"],
                                        block_type="text",
                                    )
                                )

                        if ocr_err:
                            ocr_error = ocr_err
                    except Exception as e:
                        ocr_error = f"OCR invocation error: {str(e)}"

                pages.append(
                    DocumentPage(
                        page_number=page_num,
                        text=final_text,
                        char_count=len(final_text),
                        word_count=len(final_text.split()),
                        dimensions=dim,
                        blocks=blocks,
                        tables=tables,
                        ocr_applied=ocr_applied,
                        ocr_confidence=ocr_conf,
                        ocr_error=ocr_error,
                    )
                )
        finally:
            doc.close()

        return pages

    def _parse_docx(self, file_path: Path) -> List[DocumentPage]:
        """Parses DOCX document preserving headings, paragraphs, and structured tables."""
        doc = docx.Document(file_path)
        full_text_parts: List[str] = []
        blocks: List[TextBlock] = []
        tables: List[ExtractedTable] = []

        # Extract paragraphs
        for idx, p in enumerate(doc.paragraphs):
            p_text = p.text.strip()
            if p_text:
                full_text_parts.append(p_text)
                blocks.append(
                    TextBlock(
                        block_index=idx,
                        bbox=[],
                        text=p_text,
                        block_type="heading" if p.style and "heading" in p.style.name.lower() else "text",
                    )
                )

        # Extract tables
        for t_idx, table in enumerate(doc.tables):
            table_rows: List[List[str]] = []
            for row in table.rows:
                table_rows.append([cell.text.strip() for cell in row.cells])

            if table_rows:
                headers = table_rows[0]
                rows = table_rows[1:] if len(table_rows) > 1 else []
                tables.append(
                    ExtractedTable(
                        table_index=t_idx + 1,
                        num_rows=len(table_rows),
                        num_cols=len(headers),
                        headers=headers,
                        rows=rows,
                    )
                )
                # Append table markdown representation to text
                full_text_parts.append(f"\n[Table {t_idx + 1}]:\n| " + " | ".join(headers) + " |")
                for r in rows[:10]:
                    full_text_parts.append("| " + " | ".join(r) + " |")

        joined_text = "\n\n".join(full_text_parts)
        return [
            DocumentPage(
                page_number=1,
                text=joined_text,
                char_count=len(joined_text),
                word_count=len(joined_text.split()),
                dimensions={"width": 612.0, "height": 792.0},
                blocks=blocks,
                tables=tables,
                ocr_applied=False,
            )
        ]

    def _parse_xlsx(self, file_path: Path) -> List[DocumentPage]:
        """Parses multi-sheet XLSX documents into structured tables and formatted text per sheet."""
        wb = openpyxl.load_workbook(file_path, data_only=True)
        pages: List[DocumentPage] = []

        try:
            for s_idx, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                all_rows = list(sheet.iter_rows(values_only=True))

                # Filter out completely empty rows
                non_empty_rows = [
                    [str(c).strip() if c is not None else "" for c in row]
                    for row in all_rows
                    if any(c is not None and str(c).strip() for c in row)
                ]

                tables: List[ExtractedTable] = []
                blocks: List[TextBlock] = []
                text_lines: List[str] = [f"=== Sheet: {sheet_name} ==="]

                if non_empty_rows:
                    headers = non_empty_rows[0]
                    rows = non_empty_rows[1:]
                    tables.append(
                        ExtractedTable(
                            table_index=1,
                            num_rows=len(non_empty_rows),
                            num_cols=len(headers),
                            headers=headers,
                            rows=rows,
                        )
                    )

                    text_lines.append("| " + " | ".join(headers) + " |")
                    text_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                    for r in rows:
                        text_lines.append("| " + " | ".join(r) + " |")

                    blocks.append(
                        TextBlock(
                            block_index=1,
                            bbox=[],
                            text=f"Worksheet: {sheet_name} ({len(non_empty_rows)} rows)",
                            block_type="table",
                        )
                    )

                sheet_text = "\n".join(text_lines)
                pages.append(
                    DocumentPage(
                        page_number=s_idx + 1,
                        text=sheet_text,
                        char_count=len(sheet_text),
                        word_count=len(sheet_text.split()),
                        dimensions={"width": 800.0, "height": 1000.0},
                        blocks=blocks,
                        tables=tables,
                        ocr_applied=False,
                    )
                )
        finally:
            wb.close()

        return pages

    def _parse_image(self, file_path: Path) -> List[DocumentPage]:
        """Parses standalone image (PNG, TIFF, JPG) using dual OCR engine (RapidOCR / Tesseract)."""
        with Image.open(file_path) as img:
            dim = {"width": float(img.width), "height": float(img.height)}
            ocr_text, ocr_conf, ocr_err, ocr_boxes = self.ocr_service.ocr_image(img)

        blocks: List[TextBlock] = []
        if ocr_boxes:
            for b in ocr_boxes:
                blocks.append(
                    TextBlock(
                        block_index=b["index"],
                        bbox=b["bbox"],
                        text=b["text"],
                        block_type="text",
                    )
                )
        elif ocr_text:
            for b_idx, line in enumerate(ocr_text.split("\n")):
                if line.strip():
                    blocks.append(
                        TextBlock(
                            block_index=b_idx + 1,
                            bbox=[],
                            text=line.strip(),
                            block_type="text",
                        )
                    )

        return [
            DocumentPage(
                page_number=1,
                text=ocr_text or (f"[Image Document - {file_path.name}]" if not ocr_err else f"[Image Document: {ocr_err}]"),
                char_count=len(ocr_text),
                word_count=len(ocr_text.split()),
                dimensions=dim,
                blocks=blocks,
                tables=[],
                ocr_applied=bool(ocr_text),
                ocr_confidence=ocr_conf,
                ocr_error=ocr_err,
            )
        ]

    def process_document(
        self,
        file_bytes: bytes,
        original_filename: str,
        content_type: str,
        content_hash: Optional[str] = None,
    ) -> Tuple[DocumentMetadata, List[DocumentPage]]:
        """
        Executes end-to-end ingestion pipeline:
        1. Validate inputs
        2. Generate unique Document ID
        3. Save file to storage/uploads/{doc_id}/
        4. Multi-modal parse (PDF/DOCX/XLSX/Image + OCR fallback)
        5. Persist parsed pages JSON to storage/processed/{doc_id}.json
        6. Persist metadata into SQLite DocumentDatabase
        """
        start_time = time.perf_counter()
        doc_id = str(uuid.uuid4())
        safe_filename = self.sanitize_filename(original_filename)

        is_valid, error_msg, file_type = self.validate_file(
            safe_filename, content_type, len(file_bytes)
        )
        if not is_valid:
            raise ValueError(error_msg)

        # Calculate or use provided SHA-256 content hash
        computed_hash = content_hash or hashlib.sha256(file_bytes).hexdigest()

        # 1. Save raw file
        doc_dir = self.upload_dir / doc_id
        doc_dir.mkdir(parents=True, exist_ok=True)
        raw_file_path = doc_dir / safe_filename
        raw_file_path.write_bytes(file_bytes)

        # 2. Parse by format
        pages: List[DocumentPage] = []
        parse_error: Optional[str] = None

        try:
            if file_type == "pdf":
                pages = self._parse_pdf(raw_file_path)
            elif file_type == "docx":
                pages = self._parse_docx(raw_file_path)
            elif file_type == "xlsx":
                pages = self._parse_xlsx(raw_file_path)
            elif file_type in ("png", "jpg", "tiff"):
                pages = self._parse_image(raw_file_path)
            else:
                raise ValueError(f"No parser available for type '{file_type}'")
        except Exception as e:
            parse_error = f"Parsing failed: {str(e)}"
            # Create a fallback error page
            pages = [
                DocumentPage(
                    page_number=1,
                    text=f"[Error during parsing: {str(e)}]",
                    ocr_applied=False,
                    ocr_error=parse_error,
                )
            ]

        # 3. Calculate statistics
        total_pages = len(pages)
        ocr_pages_count = sum(1 for p in pages if p.ocr_applied)
        has_tables = any(len(p.tables) > 0 for p in pages)
        preview_text = pages[0].text[:300].strip() if pages and pages[0].text else None

        processing_time_ms = int((time.perf_counter() - start_time) * 1000)

        metadata = DocumentMetadata(
            id=doc_id,
            filename=safe_filename,
            file_type=file_type,
            mime_type=content_type or f"application/{file_type}",
            file_size_bytes=len(file_bytes),
            total_pages=total_pages,
            status="processed" if not parse_error else "failed",
            uploaded_at=datetime.now(timezone.utc),
            processing_time_ms=processing_time_ms,
            ocr_pages_count=ocr_pages_count,
            has_tables=has_tables,
            preview_text=preview_text,
            error_message=parse_error,
            content_hash=computed_hash,
        )

        # 4. Save processed pages JSON
        processed_file_path = self.processed_dir / f"{doc_id}.json"
        processed_payload = {
            "metadata": metadata.model_dump(mode="json"),
            "pages": [p.model_dump(mode="json") for p in pages],
        }
        processed_file_path.write_text(json.dumps(processed_payload, indent=2), encoding="utf-8")

        # 5. Persist to SQLite
        doc_db.save_document(
            metadata=metadata,
            stored_path=str(raw_file_path),
            processed_path=str(processed_file_path),
        )

        return metadata, pages

    def get_document_pages(self, doc_id: str) -> Optional[List[DocumentPage]]:
        """Loads detailed pages list for a document from storage/processed/{doc_id}.json."""
        processed_file = self.processed_dir / f"{doc_id}.json"
        if not processed_file.exists():
            return None

        try:
            data = json.loads(processed_file.read_text(encoding="utf-8"))
            return [DocumentPage(**p) for p in data.get("pages", [])]
        except Exception:
            return None


document_service = DocumentIngestionService()
